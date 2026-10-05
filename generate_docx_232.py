import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    tblBorders = OxmlElement('w:tblBorders')
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), val)
        border.set(qn('w:sz'), sz)
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), color)
        tblBorders.append(border)
    tblPr.append(tblBorders)

def add_code_listing(doc, title, code_text):
    p_title = doc.add_paragraph()
    r_title = p_title.add_run(title)
    r_title.font.name = 'Calibri'
    r_title.font.size = Pt(11)
    r_title.font.bold = True
    p_title.paragraph_format.space_before = Pt(10)
    p_title.paragraph_format.space_after = Pt(4)

    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F8F9FA")
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for b in ['top', 'left', 'bottom', 'right']:
        node = OxmlElement(f'w:{b}')
        node.set(qn('w:val'), 'single')
        node.set(qn('w:sz'), '6')
        node.set(qn('w:color'), '005691')
        tcBorders.append(node)
    tcPr.append(tcBorders)

    p_code = cell.paragraphs[0]
    p_code.paragraph_format.space_before = Pt(2)
    p_code.paragraph_format.space_after = Pt(2)
    p_code.paragraph_format.line_spacing = 1.05
    r_code = p_code.add_run(code_text)
    r_code.font.name = 'Consolas'
    r_code.font.size = Pt(8.5)
    r_code.font.color.rgb = RGBColor(0x1F, 0x23, 0x28)

    p_sep = doc.add_paragraph()
    p_sep.paragraph_format.space_before = Pt(2)
    p_sep.paragraph_format.space_after = Pt(4)

def generate_report():
    doc = docx.Document()

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Document Header Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("Experiment – 2.3.2")
    r_title.font.name = 'Calibri'
    r_title.font.size = Pt(16)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
    p_title.paragraph_format.space_after = Pt(12)

    # Student Info Table
    info_table = doc.add_table(rows=3, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_table.autofit = False
    set_table_borders(info_table, color="003366", sz="6")

    info_data = [
        [("Student Name: ", True), ("Himanshu Goyal", False), ("UID: ", True), ("24BDA70369", False)],
        [("Branch: ", True), ("CSE (Data Science)", False), ("Section/Group: ", True), ("24 BDS-4(B)", False)],
        [("Course: ", True), ("Full Stack II (24CSP-337)", False), ("Date of Performance: ", True), ("23-09-2026", False)]
    ]

    col_widths = [Inches(3.25), Inches(3.25)]
    for row_idx, row_content in enumerate(info_data):
        row = info_table.rows[row_idx]
        for col_idx in range(2):
            cell = row.cells[col_idx]
            cell.width = col_widths[col_idx]
            set_cell_background(cell, "F2F5F8" if row_idx % 2 == 0 else "FFFFFF")
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            
            lbl_text, lbl_bold = row_content[col_idx * 2]
            val_text, val_bold = row_content[col_idx * 2 + 1]

            r1 = p.add_run(lbl_text)
            r1.font.name = 'Calibri'
            r1.font.size = Pt(10)
            r1.font.bold = lbl_bold
            r1.font.color.rgb = RGBColor(0x00, 0x33, 0x66)

            r2 = p.add_run(val_text)
            r2.font.name = 'Calibri'
            r2.font.size = Pt(10)
            r2.font.bold = val_bold

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(6)
    p_spacer.paragraph_format.space_after = Pt(6)

    def add_section_heading(num_str, title_str):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r_num = p.add_run(f"{num_str}. {title_str}:")
        r_num.font.name = 'Calibri'
        r_num.font.size = Pt(12)
        r_num.font.bold = True
        r_num.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
        return p

    def add_body_p(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = p.add_run(text)
        r.font.name = 'Calibri'
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(0x24, 0x29, 0x2F)
        return p

    def add_bullet(items):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        r_bullet = p.add_run("•  ")
        r_bullet.font.name = 'Calibri'
        r_bullet.font.size = Pt(11)
        r_bullet.font.bold = True
        r_bullet.font.color.rgb = RGBColor(0x00, 0x33, 0x66)

        for text, bold, italic in items:
            r = p.add_run(text)
            r.font.name = 'Calibri'
            r.font.size = Pt(11)
            r.font.bold = bold
            r.font.italic = italic
            r.font.color.rgb = RGBColor(0x24, 0x29, 0x2F)
        return p

    # 1. Aim
    add_section_heading("1", "Aim")
    add_body_p("To secure sensitive data using encryption techniques and manage token lifecycle strategies.")

    # 2. Software Required
    add_section_heading("2", "Software Required")
    add_bullet([("Programming Language: ", True, False), ("Java (JDK 17 or higher) with Java Cryptography Architecture (JCA)", False, False)])
    add_bullet([("Framework: ", True, False), ("Spring Boot 3.x / 4.x (Spring Security, Spring Data JPA)", False, False)])
    add_bullet([("Cryptographic Algorithms: ", True, False), ("AES-256-GCM (Galois/Counter Mode) with 96-bit IV and 128-bit authentication tag", False, False)])
    add_bullet([("IDE / Environment: ", True, False), ("Visual Studio Code with Java Extension Pack / IntelliJ IDEA", False, False)])
    add_bullet([("API Testing Tool: ", True, False), ("Postman / Windows PowerShell Invoke-RestMethod", False, False)])

    # 3. Objectives
    add_section_heading("3", "Objectives")
    add_bullet([("To understand cryptographic encryption principles: ", True, False), ("authenticated symmetric encryption, initialization vectors, tampering prevention.", False, False)])
    add_bullet([("To implement AES-256-GCM encryption: ", True, False), ("safeguarding sensitive third-party OAuth client secrets and access tokens at rest.", False, False)])
    add_bullet([("To manage token lifecycles: ", True, False), ("issuing short-lived access tokens alongside persistent, revocable refresh tokens.", False, False)])
    add_bullet([("To enforce refresh token rotation: ", True, False), ("invalidating old refresh tokens upon use and detecting replay attacks to revoke compromised sessions.", False, False)])
    add_bullet([("Course Outcomes Mapped: ", True, False), ("CO2 - BT2, CO3 - BT3, CO4 - BT4", False, True)])

    # 4. Theory
    add_section_heading("4", "Theory")
    add_body_p(
        "Storing sensitive credentials (such as third-party OAuth secrets, Facebook Graph API keys, and Twitter bearer tokens) in plaintext "
        "within relational databases creates severe vulnerability risks if database snapshots, backups, or replication logs are intercepted. "
        "Enterprise compliance frameworks (PCI-DSS, SOC 2, HIPAA) mandate encryption at rest using authenticated symmetric algorithms.\n\n"
        "Advanced Encryption Standard (AES) operated in Galois/Counter Mode (AES-256-GCM) is the industry benchmark. Unlike legacy CBC or ECB modes, "
        "GCM provides both confidentiality and built-in cryptographic authentication. Each encryption operation generates a unique, cryptographically "
        "random 96-bit Initialization Vector (IV) and outputs a 128-bit authentication tag. Any unauthorized bit modification or data tampering "
        "at rest causes immediate decryption failure with an AEADBadTagException.\n\n"
        "Simultaneously, securing user sessions requires strict token lifecycle management. Short-lived access tokens (e.g., 15 minutes) limit the "
        "window of opportunity if intercepted. Long-lived refresh tokens (e.g., 7 days) allow seamless session renewal without re-prompting users "
        "for credentials. Under Refresh Token Rotation, each refresh request invalidates the submitted token and issues a completely new token pair. "
        "If an attacker attempts to replay an already-consumed refresh token, the server detects the reuse condition, flags a potential breach, "
        "and immediately invalidates all active tokens associated with that user's session family."
    )

    # 5. Procedure / Algorithm
    add_section_heading("5", "Procedure / Algorithm")
    add_bullet([("Step 1: ", True, False), ("Implement AesEncryptionService utilizing Java JCA (Cipher.getInstance('AES/GCM/NoPadding')) with 256-bit symmetric keys.", False, False)])
    add_bullet([("Step 2: ", True, False), ("Generate random 12-byte (96-bit) IV for each encryption call; prepend IV to ciphertext and encode in Base64.", False, False)])
    add_bullet([("Step 3: ", True, False), ("Create OAuthCredential entity with encrypted clientSecret and accessToken fields.", False, False)])
    add_bullet([("Step 4: ", True, False), ("Develop OAuthCredentialService encrypting credentials before database persist and masking secrets upon retrieval.", False, False)])
    add_bullet([("Step 5: ", True, False), ("Expose /api/v1/credentials and /api/v1/credentials/{serviceName}/raw-audit endpoints verifying zero plaintext exposure in database rows.", False, False)])
    add_bullet([("Step 6: ", True, False), ("Create RefreshToken entity tracking token UUID, user link, expiry date, revoked flag, and family lineage.", False, False)])
    add_bullet([("Step 7: ", True, False), ("Implement TokenLifecycleService managing /api/v1/auth/refresh, enforcing single-use token rotation.", False, False)])
    add_bullet([("Step 8: ", True, False), ("Implement replay attack prevention: if an already-revoked refresh token is presented, revoke all active tokens for that user.", False, False)])
    add_bullet([("Step 9: ", True, False), ("Verify encryption at rest, token rotation, and replay mitigation via 4 automated integration tests in EncryptionAndTokenLifecycleTest.", False, False)])

    # 6. Code and Output
    add_section_heading("6", "Code and Output")

    code_aes = """@Service
public class AesEncryptionService {

    private static final String ALGORITHM = "AES/GCM/NoPadding";
    private static final int GCM_TAG_LENGTH_BITS = 128;
    private static final int IV_LENGTH_BYTES = 12;

    @Value("${app.security.encryption.key}")
    private String base64SecretKey;

    public String encrypt(String plaintext) {
        if (plaintext == null) return null;
        try {
            byte[] iv = new byte[IV_LENGTH_BYTES];
            SecureRandom.getInstanceStrong().nextBytes(iv);

            Cipher cipher = Cipher.getInstance(ALGORITHM);
            cipher.init(Cipher.ENCRYPT_MODE, getSecretKeySpec(), new GCMParameterSpec(GCM_TAG_LENGTH_BITS, iv));

            byte[] ciphertext = cipher.doFinal(plaintext.getBytes(StandardCharsets.UTF_8));
            byte[] payload = ByteBuffer.allocate(iv.length + ciphertext.length)
                    .put(iv).put(ciphertext).array();

            return Base64.getEncoder().encodeToString(payload);
        } catch (Exception ex) {
            throw new IllegalStateException("AES-256-GCM encryption failed", ex);
        }
    }
}"""
    add_code_listing(doc, "Listing 1: AES-256-GCM Authenticated Encryption Service (AesEncryptionService)", code_aes)

    code_refresh = """@Service
public class TokenLifecycleService {

    @Transactional
    public AuthResponse refreshAccessToken(String rawRefreshToken) {
        RefreshToken token = refreshTokenRepository.findByToken(rawRefreshToken)
                .orElseThrow(() -> new BadRequestException("Invalid refresh token."));

        // Replay Attack Detection: Token reuse triggers nuclear session revocation
        if (token.isRevoked()) {
            log.warn("[Security Alert] Revoked refresh token reuse detected for user: {}. Revoking all sessions.",
                    token.getUser().getUsername());
            refreshTokenRepository.revokeAllUserTokens(token.getUser());
            throw new BadRequestException("Security Breach Detected: Token reuse attempt. All active sessions invalidated.");
        }

        // Token Rotation: Invalidate current token and issue fresh token pair
        token.setRevoked(true);
        refreshTokenRepository.save(token);

        RefreshToken newRefreshToken = createRefreshToken(token.getUser());
        String newAccessToken = jwtService.generateToken(token.getUser());

        return new AuthResponse(newAccessToken, newRefreshToken.getToken(), 900, "Bearer");
    }
}"""
    add_code_listing(doc, "Listing 2: Token Lifecycle Service with Refresh Token Rotation & Replay Attack Mitigation", code_refresh)

    # Screenshots / Execution Outputs
    img_dir = r"C:\Users\himan\Downloads\social-post-backend\social-post-backend\docs\screenshots"
    img1_path = os.path.join(img_dir, "ss_exp232_aes_encryption.png")
    img2_path = os.path.join(img_dir, "ss_exp232_token_refresh.png")

    if os.path.exists(img1_path):
        p_img_title1 = doc.add_paragraph()
        r_it1 = p_img_title1.add_run("Output 1: AES-256-GCM Credential Storage & Raw Database Column Audit (Proving Zero Plaintext Exposure)")
        r_it1.font.name = 'Calibri'
        r_it1.font.bold = True
        p_img1 = doc.add_paragraph()
        p_img1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img1.add_run().add_picture(img1_path, width=Inches(6.5))

    if os.path.exists(img2_path):
        p_img_title2 = doc.add_paragraph()
        r_it2 = p_img_title2.add_run("Output 2: Refresh Token Rotation & Replay Attack Mitigation (Automated Immediate Revocation of Reused Tokens)")
        r_it2.font.name = 'Calibri'
        r_it2.font.bold = True
        p_img2 = doc.add_paragraph()
        p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img2.add_run().add_picture(img2_path, width=Inches(6.5))

    # 7. Conclusion
    add_section_heading("7", "Conclusion")
    add_body_p(
        "Enterprise-grade data security and session protection were successfully implemented. Sensitive OAuth credentials and secrets are "
        "cryptographically guarded at rest using AES-256-GCM authenticated encryption with dynamic 96-bit initialization vectors and 128-bit authentication tags, "
        "guaranteeing both privacy and tamper detection. Direct database audits proved zero plaintext exposure. "
        "Token lifecycle management demonstrated robust session defense via short-lived JWT access tokens and single-use refresh token rotation. "
        "Replay attack mitigation logic successfully detected duplicate token presentations and invalidated all family sessions, achieving 100% test verification."
    )

    # 8. Learning Outcomes
    add_section_heading("8", "Learning Outcomes")
    add_bullet([("CO2 - BT2: ", True, False), ("Understood authenticated encryption mechanisms (AES-256-GCM), initialization vectors, authentication tags, and tamper resistance.", False, False)])
    add_bullet([("CO3 - BT3: ", True, False), ("Engineered encryption services and persistent JPA entities safeguarding sensitive credentials and API tokens at rest.", False, False)])
    add_bullet([("CO4 - BT4: ", True, False), ("Evaluated session security architectures, implementing refresh token rotation and automated replay attack revocation.", False, False)])
    add_bullet([("CO4 - BT4: ", True, False), ("Conducted raw database column audits, confirming zero plaintext exposure across persistent relational storage.", False, False)])

    out_repo = r"C:\Users\himan\Downloads\social-post-backend\social-post-backend\docs\lab-reports\FSD 2.3.2.docx"
    doc.save(out_repo)
    print(f"Generated repo report: {out_repo}")

    try:
        out_dl = r"C:\Users\himan\Downloads\FSD 2.3.2.docx"
        doc.save(out_dl)
        print(f"Generated Downloads report: {out_dl}")
    except Exception as e:
        print(f"Could not overwrite {out_dl}: {e}")

if __name__ == "__main__":
    generate_report()
