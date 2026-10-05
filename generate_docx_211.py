import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    tblBorders = OxmlElement('w:tblBorders')
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), val)
        border.set(qn('w:sz'), sz)
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), color)
        tblBorders.append(border)
    tblPr.append(tblBorders)

def add_code_listing(doc, title, code_text):
    p_title = doc.add_paragraph()
    r_title = p_title.add_run(title)
    r_title.font.name = 'Calibri'
    r_title.font.size = Pt(11)
    r_title.font.bold = True
    p_title.paragraph_format.space_before = Pt(10)
    p_title.paragraph_format.space_after = Pt(4)

    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F8F9FA")
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for b in ['top', 'left', 'bottom', 'right']:
        node = OxmlElement(f'w:{b}')
        node.set(qn('w:val'), 'single')
        node.set(qn('w:sz'), '6')
        node.set(qn('w:color'), '005691')
        tcBorders.append(node)
    tcPr.append(tcBorders)

    p_code = cell.paragraphs[0]
    p_code.paragraph_format.space_before = Pt(2)
    p_code.paragraph_format.space_after = Pt(2)
    p_code.paragraph_format.line_spacing = 1.05
    r_code = p_code.add_run(code_text)
    r_code.font.name = 'Consolas'
    r_code.font.size = Pt(8.5)
    r_code.font.color.rgb = RGBColor(0x1F, 0x23, 0x28)

    p_sep = doc.add_paragraph()
    p_sep.paragraph_format.space_before = Pt(2)
    p_sep.paragraph_format.space_after = Pt(4)

def generate_report():
    doc = docx.Document()

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Document Header Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("Experiment – 2.1.1")
    r_title.font.name = 'Calibri'
    r_title.font.size = Pt(16)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
    p_title.paragraph_format.space_after = Pt(12)

    # Student Info Table
    info_table = doc.add_table(rows=3, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_table.autofit = False
    set_table_borders(info_table, color="003366", sz="6")

    info_data = [
        [("Student Name: ", True), ("Himanshu Goyal", False), ("UID: ", True), ("24BDA70369", False)],
        [("Branch: ", True), ("CSE (Data Science)", False), ("Section/Group: ", True), ("24 BDS-4(B)", False)],
        [("Course: ", True), ("Full Stack II (24CSP-337)", False), ("Date of Performance: ", True), ("09-09-2026", False)]
    ]

    col_widths = [Inches(3.25), Inches(3.25)]
    for row_idx, row_content in enumerate(info_data):
        row = info_table.rows[row_idx]
        for col_idx in range(2):
            cell = row.cells[col_idx]
            cell.width = col_widths[col_idx]
            set_cell_background(cell, "F2F5F8" if row_idx % 2 == 0 else "FFFFFF")
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            
            lbl_text, lbl_bold = row_content[col_idx * 2]
            val_text, val_bold = row_content[col_idx * 2 + 1]

            r1 = p.add_run(lbl_text)
            r1.font.name = 'Calibri'
            r1.font.size = Pt(10)
            r1.font.bold = lbl_bold
            r1.font.color.rgb = RGBColor(0x00, 0x33, 0x66)

            r2 = p.add_run(val_text)
            r2.font.name = 'Calibri'
            r2.font.size = Pt(10)
            r2.font.bold = val_bold

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(6)
    p_spacer.paragraph_format.space_after = Pt(6)

    def add_section_heading(num_str, title_str):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r_num = p.add_run(f"{num_str}. {title_str}:")
        r_num.font.name = 'Calibri'
        r_num.font.size = Pt(12)
        r_num.font.bold = True
        r_num.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
        return p

    def add_body_p(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = p.add_run(text)
        r.font.name = 'Calibri'
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(0x24, 0x29, 0x2F)
        return p

    def add_bullet(items):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        r_bullet = p.add_run("•  ")
        r_bullet.font.name = 'Calibri'
        r_bullet.font.size = Pt(11)
        r_bullet.font.bold = True
        r_bullet.font.color.rgb = RGBColor(0x00, 0x33, 0x66)

        for text, bold, italic in items:
            r = p.add_run(text)
            r.font.name = 'Calibri'
            r.font.size = Pt(11)
            r.font.bold = bold
            r.font.italic = italic
            r.font.color.rgb = RGBColor(0x24, 0x29, 0x2F)
        return p

    # 1. Aim
    add_section_heading("1", "Aim")
    add_body_p("To design and implement RESTful APIs using Spring Boot with proper validation, standardized responses, and scalable architecture.")

    # 2. Software Required
    add_section_heading("2", "Software Required")
    add_bullet([("Programming Language: ", True, False), ("Java (JDK 17 or higher)", False, False)])
    add_bullet([("Framework: ", True, False), ("Spring Boot 3.x / 4.x (Web, Validation, Spring Data JPA)", False, False)])
    add_bullet([("IDE / Editor: ", True, False), ("Visual Studio Code with Java Extension Pack / IntelliJ IDEA", False, False)])
    add_bullet([("API Testing Tool: ", True, False), ("Postman / Windows PowerShell Invoke-RestMethod", False, False)])
    add_bullet([("Relational Database: ", True, False), ("In-Memory H2 Database Engine / PostgreSQL", False, False)])

    # 3. Objectives
    add_section_heading("3", "Objectives")
    add_bullet([("To understand REST API design principles: ", True, False), ("stateless communication, resource-oriented URI endpoints, standard HTTP verbs.", False, False)])
    add_bullet([("To implement CRUD APIs using Spring Boot: ", True, False), ("layered architecture with Controller, Service, and Repository layers.", False, False)])
    add_bullet([("To enforce consistent request-response structures: ", True, False), ("encapsulating status codes, messages, and payloads in ApiResponse<T>.", False, False)])
    add_bullet([("To apply validation using Bean Validation: ", True, False), ("using @NotBlank, @Size, and @NotNull annotations on request DTOs.", False, False)])
    add_bullet([("To enable secure cross-origin communication: ", True, False), ("configuring CORS mappings for decoupled frontend-backend architectures.", False, False)])
    add_bullet([("Course Outcomes Mapped: ", True, False), ("CO1 - BT1, CO3 - BT3", False, True)])

    # 4. Theory
    add_section_heading("4", "Theory")
    add_body_p(
        "REST (Representational State Transfer) is an architectural style used for designing scalable and loosely coupled web services. "
        "In RESTful systems, data is treated as resources identified by uniform resource identifiers (URIs) and accessed through standard "
        "HTTP methods such as GET (read), POST (create), PUT (update), and DELETE (remove). A well-designed API must adhere to stateless communication, "
        "meaning every request from client to server contains all necessary context for processing without server-side session reliance.\n\n"
        "Spring Boot simplifies enterprise REST API engineering through an established layered architecture:\n"
        "1. Controller Layer: Exposes HTTP endpoints, consumes request bodies, performs input binding, and returns standardized response entities.\n"
        "2. Service Layer: Contains business domain logic, transactional boundaries, validation rules, and entity-DTO transformation.\n"
        "3. Repository Layer: Leverages Spring Data JPA to execute database CRUD operations without boilerplate SQL.\n\n"
        "To guarantee data integrity, Bean Validation is applied to request DTOs. Annotations such as @NotNull, @NotBlank, and @Size validate "
        "incoming payloads before reaching business logic, triggering automated 400 Bad Request responses when violated. Standardizing responses "
        "using an envelope structure (ApiResponse<T>) ensures predictable contracts for frontend clients across both successful operations and errors. "
        "Cross-Origin Resource Sharing (CORS) is explicitly configured to enable secure communication between decoupled clients and backend services."
    )

    # 5. Procedure / Algorithm
    add_section_heading("5", "Procedure / Algorithm")
    add_bullet([("Step 1: ", True, False), ("Initialize Spring Boot project structure with Spring Web, Validation, and JPA starter dependencies.", False, False)])
    add_bullet([("Step 2: ", True, False), ("Define domain model (Post) with JPA annotations (@Entity, @Table, @Id, @GeneratedValue, @Column).", False, False)])
    add_bullet([("Step 3: ", True, False), ("Create PostRepository interface extending JpaRepository<Post, Long> for automated CRUD queries.", False, False)])
    add_bullet([("Step 4: ", True, False), ("Implement PostService and PostServiceImpl containing business logic for create, read, update, and delete.", False, False)])
    add_bullet([("Step 5: ", True, False), ("Expose RESTful endpoints under /api/v1/posts in PostController handling GET, POST, PUT, DELETE.", False, False)])
    add_bullet([("Step 6: ", True, False), ("Enforce Bean Validation on CreatePostRequest and UpdatePostRequest DTOs using @Valid in controller methods.", False, False)])
    add_bullet([("Step 7: ", True, False), ("Wrap all endpoints in generic ApiResponse<T> envelope returning success, message, data, and ISO-8601 timestamp.", False, False)])
    add_bullet([("Step 8: ", True, False), ("Configure WebMvcConfigurer CORS bean to permit frontend origins, allowed headers, and HTTP methods.", False, False)])
    add_bullet([("Step 9: ", True, False), ("Verify CRUD operations and validation failure boundaries using automated MockMvc integration tests and PowerShell.", False, False)])

    # 6. Code and Output
    add_section_heading("6", "Code and Output")

    code_controller = """@RestController
@RequestMapping("/api/v1/posts")
public class PostController {

    private final PostService postService;

    public PostController(PostService postService) {
        this.postService = postService;
    }

    @PostMapping
    public ResponseEntity<ApiResponse<PostResponse>> createPost(@Valid @RequestBody CreatePostRequest request) {
        PostResponse response = postService.createPost(request);
        return ResponseEntity.status(HttpStatus.CREATED)
                .body(ApiResponse.created(response, "Post created successfully"));
    }

    @GetMapping("/{id}")
    public ResponseEntity<ApiResponse<PostResponse>> getPostById(@PathVariable Long id) {
        PostResponse response = postService.getPostById(id);
        return ResponseEntity.ok(ApiResponse.ok(response));
    }

    @PutMapping("/{id}")
    public ResponseEntity<ApiResponse<PostResponse>> updatePost(
            @PathVariable Long id,
            @Valid @RequestBody UpdatePostRequest request) {
        PostResponse response = postService.updatePost(id, request);
        return ResponseEntity.ok(ApiResponse.ok(response, "Post updated successfully"));
    }

    @DeleteMapping("/{id}")
    public ResponseEntity<ApiResponse<Void>> deletePost(@PathVariable Long id) {
        postService.deletePost(id);
        return ResponseEntity.ok(ApiResponse.ok(null, "Post deleted successfully"));
    }
}"""
    add_code_listing(doc, "Listing 1: Layered Architecture – PostController REST Endpoints with Validation", code_controller)

    code_dto = """public class CreatePostRequest {

    @NotBlank(message = "Title is required and cannot be blank")
    @Size(min = 3, max = 150, message = "Title must be between 3 and 150 characters")
    private String title;

    @NotBlank(message = "Content is required and cannot be blank")
    @Size(min = 10, message = "Content must be at least 10 characters long")
    private String content;

    @NotBlank(message = "Author is required")
    private String author;

    // Getters, Setters, and Constructors
}

public class ApiResponse<T> {
    private boolean success;
    private int status;
    private String message;
    private T data;
    private Instant timestamp;

    public static <T> ApiResponse<T> ok(T data, String message) {
        return new ApiResponse<>(true, 200, message, data, Instant.now());
    }

    public static <T> ApiResponse<T> created(T data, String message) {
        return new ApiResponse<>(true, 201, message, data, Instant.now());
    }
}"""
    add_code_listing(doc, "Listing 2: Bean Validation DTO and Standardized Envelope (ApiResponse<T>)", code_dto)

    code_cors = """@Configuration
public class WebConfig implements WebMvcConfigurer {

    @Override
    public void addCorsMappings(CorsRegistry registry) {
        registry.addMapping("/api/**")
                .allowedOrigins("http://localhost:3000", "http://localhost:5173", "https://*.onrender.com")
                .allowedMethods("GET", "POST", "PUT", "PATCH", "DELETE", "OPTIONS")
                .allowedHeaders("*")
                .exposedHeaders("X-Correlation-ID", "Authorization")
                .allowCredentials(true)
                .maxAge(3600);
    }
}"""
    add_code_listing(doc, "Listing 3: CORS Configuration for Secure Cross-Origin Frontend Integration", code_cors)

    # Screenshots / Execution Outputs
    img_dir = r"C:\Users\himan\Downloads\social-post-backend\social-post-backend\docs\screenshots"
    img1_path = os.path.join(img_dir, "ss_exp211_img1.png")
    img2_path = os.path.join(img_dir, "ss_exp211_img2.png")

    if os.path.exists(img1_path):
        p_img_title1 = doc.add_paragraph()
        r_it1 = p_img_title1.add_run("Output 1: Create Post (201 Created) and Get All Posts with Standardized Envelope")
        r_it1.font.name = 'Calibri'
        r_it1.font.bold = True
        p_img1 = doc.add_paragraph()
        p_img1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img1.add_run().add_picture(img1_path, width=Inches(6.5))

    if os.path.exists(img2_path):
        p_img_title2 = doc.add_paragraph()
        r_it2 = p_img_title2.add_run("Output 2: Get One, Update, Delete, and Bean Validation Error Response (400 Bad Request)")
        r_it2.font.name = 'Calibri'
        r_it2.font.bold = True
        p_img2 = doc.add_paragraph()
        p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img2.add_run().add_picture(img2_path, width=Inches(6.5))

    # 7. Conclusion
    add_section_heading("7", "Conclusion")
    add_body_p(
        "A robust and scalable RESTful API was successfully designed, developed, and verified using Spring Boot's layered architecture "
        "(Controller-Service-Repository). All five fundamental CRUD operations (Create, Read One, Read All, Update, Delete) were implemented "
        "for the Post entity. Incoming client payloads are defensively validated using Bean Validation annotations (@NotBlank, @Size), strictly "
        "rejecting malformed data before reaching business logic. Every API response adheres to a predictable, standardized ApiResponse<T> "
        "envelope with HTTP status codes and ISO-8601 timestamps, ensuring flawless frontend contract compatibility. Secure cross-origin communication "
        "was verified through WebMvc CORS configuration, successfully meeting all architectural and design objectives."
    )

    # 8. Learning Outcomes
    add_section_heading("8", "Learning Outcomes")
    add_bullet([("CO1 - BT1: ", True, False), ("Learned fundamental REST API design principles including resource URI naming, stateless HTTP communication, and standard verb semantics.", False, False)])
    add_bullet([("CO3 - BT3: ", True, False), ("Applied Bean Validation constraints to enforce request data integrity before processing by business domain services.", False, False)])
    add_bullet([("CO3 - BT3: ", True, False), ("Engineered standardized generic API response envelopes (ApiResponse<T>) for predictable frontend consumption.", False, False)])
    add_bullet([("CO3 - BT3: ", True, False), ("Configured Cross-Origin Resource Sharing (CORS) rules to enable secure, decoupled frontend-backend communication across diverse web origins.", False, False)])

    # Save to both repo docs/lab-reports and Downloads
    out_repo = r"C:\Users\himan\Downloads\social-post-backend\social-post-backend\docs\lab-reports\FSD 2.1.1.docx"
    doc.save(out_repo)
    print(f"Generated repo report: {out_repo}")

    try:
        out_dl = r"C:\Users\himan\Downloads\FSD 2.1.1.docx"
        doc.save(out_dl)
        print(f"Generated Downloads report: {out_dl}")
    except Exception as e:
        print(f"Could not overwrite {out_dl}: {e}")

if __name__ == "__main__":
    generate_report()
