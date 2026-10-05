import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    tblBorders = OxmlElement('w:tblBorders')
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), val)
        border.set(qn('w:sz'), sz)
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), color)
        tblBorders.append(border)
    tblPr.append(tblBorders)

def add_code_listing(doc, title, code_text):
    p_title = doc.add_paragraph()
    r_title = p_title.add_run(title)
    r_title.font.name = 'Calibri'
    r_title.font.size = Pt(11)
    r_title.font.bold = True
    p_title.paragraph_format.space_before = Pt(10)
    p_title.paragraph_format.space_after = Pt(4)

    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F8F9FA")
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for b in ['top', 'left', 'bottom', 'right']:
        node = OxmlElement(f'w:{b}')
        node.set(qn('w:val'), 'single')
        node.set(qn('w:sz'), '6')
        node.set(qn('w:color'), '005691')
        tcBorders.append(node)
    tcPr.append(tcBorders)

    p_code = cell.paragraphs[0]
    p_code.paragraph_format.space_before = Pt(2)
    p_code.paragraph_format.space_after = Pt(2)
    p_code.paragraph_format.line_spacing = 1.05
    r_code = p_code.add_run(code_text)
    r_code.font.name = 'Consolas'
    r_code.font.size = Pt(8.5)
    r_code.font.color.rgb = RGBColor(0x1F, 0x23, 0x28)

    p_sep = doc.add_paragraph()
    p_sep.paragraph_format.space_before = Pt(2)
    p_sep.paragraph_format.space_after = Pt(4)

def generate_report():
    doc = docx.Document()

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Document Header Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("Experiment – 2.3.1")
    r_title.font.name = 'Calibri'
    r_title.font.size = Pt(16)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
    p_title.paragraph_format.space_after = Pt(12)

    # Student Info Table
    info_table = doc.add_table(rows=3, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_table.autofit = False
    set_table_borders(info_table, color="003366", sz="6")

    info_data = [
        [("Student Name: ", True), ("Himanshu Goyal", False), ("UID: ", True), ("24BDA70369", False)],
        [("Branch: ", True), ("CSE (Data Science)", False), ("Section/Group: ", True), ("24 BDS-4(B)", False)],
        [("Course: ", True), ("Full Stack II (24CSP-337)", False), ("Date of Performance: ", True), ("23-09-2026", False)]
    ]

    col_widths = [Inches(3.25), Inches(3.25)]
    for row_idx, row_content in enumerate(info_data):
        row = info_table.rows[row_idx]
        for col_idx in range(2):
            cell = row.cells[col_idx]
            cell.width = col_widths[col_idx]
            set_cell_background(cell, "F2F5F8" if row_idx % 2 == 0 else "FFFFFF")
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            
            lbl_text, lbl_bold = row_content[col_idx * 2]
            val_text, val_bold = row_content[col_idx * 2 + 1]

            r1 = p.add_run(lbl_text)
            r1.font.name = 'Calibri'
            r1.font.size = Pt(10)
            r1.font.bold = lbl_bold
            r1.font.color.rgb = RGBColor(0x00, 0x33, 0x66)

            r2 = p.add_run(val_text)
            r2.font.name = 'Calibri'
            r2.font.size = Pt(10)
            r2.font.bold = val_bold

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(6)
    p_spacer.paragraph_format.space_after = Pt(6)

    def add_section_heading(num_str, title_str):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r_num = p.add_run(f"{num_str}. {title_str}:")
        r_num.font.name = 'Calibri'
        r_num.font.size = Pt(12)
        r_num.font.bold = True
        r_num.font.color.rgb = RGBColor(0x00, 0x33, 0x66)
        return p

    def add_body_p(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = p.add_run(text)
        r.font.name = 'Calibri'
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(0x24, 0x29, 0x2F)
        return p

    def add_bullet(items):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        r_bullet = p.add_run("•  ")
        r_bullet.font.name = 'Calibri'
        r_bullet.font.size = Pt(11)
        r_bullet.font.bold = True
        r_bullet.font.color.rgb = RGBColor(0x00, 0x33, 0x66)

        for text, bold, italic in items:
            r = p.add_run(text)
            r.font.name = 'Calibri'
            r.font.size = Pt(11)
            r.font.bold = bold
            r.font.italic = italic
            r.font.color.rgb = RGBColor(0x24, 0x29, 0x2F)
        return p

    # 1. Aim
    add_section_heading("1", "Aim")
    add_body_p("To implement secure backend APIs using JWT authentication and role-based authorization.")

    # 2. Software Required
    add_section_heading("2", "Software Required")
    add_bullet([("Programming Language: ", True, False), ("Java (JDK 17 or higher)", False, False)])
    add_bullet([("Framework: ", True, False), ("Spring Boot with Spring Security (spring-boot-starter-security)", False, False)])
    add_bullet([("JWT Library: ", True, False), ("JJWT (io.jsonwebtoken: jjwt-api, jjwt-impl, jjwt-jackson)", False, False)])
    add_bullet([("IDE / Environment: ", True, False), ("Visual Studio Code with Java Extension Pack / IntelliJ IDEA", False, False)])
    add_bullet([("API Testing Tool: ", True, False), ("Postman / Windows PowerShell Invoke-RestMethod", False, False)])

    # 3. Objectives
    add_section_heading("3", "Objectives")
    add_bullet([("To understand JWT-based authentication: ", True, False), ("stateless Bearer token issuance, HMAC-SHA256 signature verification.", False, False)])
    add_bullet([("To configure Spring Security for API protection: ", True, False), ("disabling CSRF, enforcing STATELESS session policy, securing endpoints.", False, False)])
    add_bullet([("To implement authentication filters: ", True, False), ("custom OncePerRequestFilter parsing Authorization header and populating SecurityContext.", False, False)])
    add_bullet([("To enforce role-based access control (RBAC): ", True, False), ("using method security @PreAuthorize('hasRole(\"ADMIN\")').", False, False)])
    add_bullet([("Course Outcomes Mapped: ", True, False), ("CO2 - BT2, CO3 - BT3, CO4 - BT4", False, True)])

    # 4. Theory
    add_section_heading("4", "Theory")
    add_body_p(
        "Modern cloud-native backends require stateless authentication to achieve effortless horizontal scalability. "
        "Traditional session-based authentication stores user session state in server memory, requiring sticky sessions or distributed session caches. "
        "JSON Web Tokens (JWT) eliminate this overhead by packaging user identity, assigned roles, and expiration timestamps into a self-contained, "
        "cryptographically signed token. The client passes this token in the 'Authorization: Bearer <token>' header on every API invocation.\n\n"
        "Spring Security intercepts all incoming requests through a customizable SecurityFilterChain. By declaring SessionCreationPolicy.STATELESS, "
        "the server never creates HTTP sessions. A custom JwtAuthenticationFilter (extending OncePerRequestFilter) extracts the Bearer token, "
        "validates its HMAC-SHA256 cryptographic signature against a secure 256-bit secret key, and constructs a UsernamePasswordAuthenticationToken "
        "populated with GrantedAuthority roles (e.g., 'ROLE_USER', 'ROLE_ADMIN').\n\n"
        "Role-Based Access Control (RBAC) restricts access to sensitive operations based on the caller's verified authorities. Method-level security "
        "enabled by @EnableMethodSecurity allows declarative annotations such as @PreAuthorize(\"hasRole('ADMIN')\"). If an unauthenticated user "
        "attempts access, Spring Security returns 401 Unauthorized; if an authenticated user with ROLE_USER attempts an ADMIN operation, "
        "it returns 403 Forbidden with a standardized ApiErrorResponse envelope."
    )

    # 5. Procedure / Algorithm
    add_section_heading("5", "Procedure / Algorithm")
    add_bullet([("Step 1: ", True, False), ("Add spring-boot-starter-security and jjwt dependencies to pom.xml.", False, False)])
    add_bullet([("Step 2: ", True, False), ("Define User entity with UserRole enum ('ROLE_USER', 'ROLE_ADMIN') and BCrypt password encryption.", False, False)])
    add_bullet([("Step 3: ", True, False), ("Implement JwtService to generate, sign, and validate HMAC-SHA256 tokens with configurable expiration (e.g., 15 minutes).", False, False)])
    add_bullet([("Step 4: ", True, False), ("Develop JwtAuthenticationFilter extending OncePerRequestFilter to intercept requests and populate SecurityContextHolder.", False, False)])
    add_bullet([("Step 5: ", True, False), ("Configure SecurityConfig bean disabling CSRF, configuring CORS, and registering JwtAuthenticationFilter before UsernamePasswordAuthenticationFilter.", False, False)])
    add_bullet([("Step 6: ", True, False), ("Enable method security with @EnableMethodSecurity(prePostEnabled = true).", False, False)])
    add_bullet([("Step 7: ", True, False), ("Develop AuthController with /register, /login, and /me endpoints returning standardized AuthResponse tokens.", False, False)])
    add_bullet([("Step 8: ", True, False), ("Develop AdminController annotated with @PreAuthorize(\"hasRole('ADMIN')\") for dashboard and purge endpoints.", False, False)])
    add_bullet([("Step 9: ", True, False), ("Verify 401 Unauthorized and 403 Forbidden enforcement using 6 automated integration tests in SecurityAndJwtIntegrationTest.", False, False)])

    # 6. Code and Output
    add_section_heading("6", "Code and Output")

    code_sec = """@Configuration
@EnableWebSecurity
@EnableMethodSecurity(prePostEnabled = true)
public class SecurityConfig {

    private final JwtAuthenticationFilter jwtAuthFilter;
    private final CustomAuthenticationEntryPoint authenticationEntryPoint;
    private final CustomAccessDeniedHandler accessDeniedHandler;

    @Bean
    public SecurityFilterChain securityFilterChain(HttpSecurity http) throws Exception {
        return http
                .csrf(AbstractHttpConfigurer::disable)
                .cors(Customizer.withDefaults())
                .sessionManagement(s -> s.sessionCreationPolicy(SessionCreationPolicy.STATELESS))
                .authorizeHttpRequests(auth -> auth
                        .requestMatchers("/api/v1/auth/**", "/h2-console/**").permitAll()
                        .requestMatchers(HttpMethod.GET, "/api/v1/posts/**").permitAll()
                        .requestMatchers("/api/v1/admin/**").hasRole("ADMIN")
                        .anyRequest().authenticated()
                )
                .exceptionHandling(ex -> ex
                        .authenticationEntryPoint(authenticationEntryPoint)
                        .accessDeniedHandler(accessDeniedHandler)
                )
                .addFilterBefore(jwtAuthFilter, UsernamePasswordAuthenticationFilter.class)
                .build();
    }
}"""
    add_code_listing(doc, "Listing 1: Stateless Security Filter Chain & RBAC Configuration (SecurityConfig)", code_sec)

    code_jwt = """@Service
public class JwtService {

    @Value("${app.security.jwt.secret}")
    private String jwtSecret;

    @Value("${app.security.jwt.expiration-seconds:900}")
    private long jwtExpirationSeconds;

    public String generateToken(UserDetails userDetails, Map<String, Object> extraClaims) {
        Instant now = Instant.now();
        return Jwts.builder()
                .claims(extraClaims)
                .subject(userDetails.getUsername())
                .issuedAt(Date.from(now))
                .expiration(Date.from(now.plusSeconds(jwtExpirationSeconds)))
                .signWith(getSigningKey(), Jwts.SIG.HS256)
                .compact();
    }

    public boolean validateToken(String token, UserDetails userDetails) {
        final String username = extractUsername(token);
        return username.equals(userDetails.getUsername()) && !isTokenExpired(token);
    }
}"""
    add_code_listing(doc, "Listing 2: Cryptographic JWT Service (HMAC-SHA256 Token Issuance & Validation)", code_jwt)

    # Screenshots / Execution Outputs
    img_dir = r"C:\Users\himan\Downloads\social-post-backend\social-post-backend\docs\screenshots"
    img1_path = os.path.join(img_dir, "ss_exp231_jwt_auth.png")
    img2_path = os.path.join(img_dir, "ss_exp231_rbac_403.png")

    if os.path.exists(img1_path):
        p_img_title1 = doc.add_paragraph()
        r_it1 = p_img_title1.add_run("Output 1: User Login (Issuing JWT Access & Refresh Tokens) and Authenticated Profile Access")
        r_it1.font.name = 'Calibri'
        r_it1.font.bold = True
        p_img1 = doc.add_paragraph()
        p_img1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img1.add_run().add_picture(img1_path, width=Inches(6.5))

    if os.path.exists(img2_path):
        p_img_title2 = doc.add_paragraph()
        r_it2 = p_img_title2.add_run("Output 2: RBAC Enforcement – Standard User Blocked (403 Forbidden) vs Admin Access Granted (200 OK)")
        r_it2.font.name = 'Calibri'
        r_it2.font.bold = True
        p_img2 = doc.add_paragraph()
        p_img2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img2.add_run().add_picture(img2_path, width=Inches(6.5))

    # 7. Conclusion
    add_section_heading("7", "Conclusion")
    add_body_p(
        "Stateless JWT authentication and Role-Based Access Control (RBAC) were successfully integrated into the Spring Boot backend. "
        "By enforcing stateless session management, eliminating CSRF vulnerability surfaces on REST endpoints, and validating cryptographically "
        "signed HMAC-SHA256 tokens, the application achieves bulletproof security and horizontal scalability. "
        "RBAC rules guarded by @PreAuthorize strictly isolated administrative capabilities, preventing unauthorized role elevation and verifying "
        "both 401 Unauthorized and 403 Forbidden failure paths through automated integration testing."
    )

    # 8. Learning Outcomes
    add_section_heading("8", "Learning Outcomes")
    add_bullet([("CO2 - BT2: ", True, False), ("Understood stateless JWT authentication principles, token structure (Header, Payload, Signature), and HMAC-SHA256 algorithms.", False, False)])
    add_bullet([("CO3 - BT3: ", True, False), ("Configured Spring Security 6/7 filter chains, disabled CSRF for stateless REST APIs, and managed SecurityContextHolder state.", False, False)])
    add_bullet([("CO3 - BT3: ", True, False), ("Implemented custom OncePerRequestFilter for Bearer token validation and GrantedAuthority role extraction.", False, False)])
    add_bullet([("CO4 - BT4: ", True, False), ("Enforced Role-Based Access Control (RBAC) using method security @PreAuthorize, distinguishing ROLE_USER and ROLE_ADMIN access rights.", False, False)])

    out_repo = r"C:\Users\himan\Downloads\social-post-backend\social-post-backend\docs\lab-reports\FSD 2.3.1.docx"
    doc.save(out_repo)
    print(f"Generated repo report: {out_repo}")

    try:
        out_dl = r"C:\Users\himan\Downloads\FSD 2.3.1.docx"
        doc.save(out_dl)
        print(f"Generated Downloads report: {out_dl}")
    except Exception as e:
        print(f"Could not overwrite {out_dl}: {e}")

if __name__ == "__main__":
    generate_report()
